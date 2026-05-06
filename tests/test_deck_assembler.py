"""Tests for :class:`shared.deck_assembler.DeckAssembler`."""

from __future__ import annotations

import io

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from shared.deck_assembler import DeckAssembler


# ---------------------------------------------------------- fixture helpers


def _name_shape(shape, name: str) -> None:
    """Set the PPTX shape ``name`` (used by python-pptx's shape index)."""
    shape.name = name


def _set_slide_name(slide, name: str) -> None:
    """Set the cSld@name attribute used by DeckAssembler's slide index."""
    cSld = slide.element.find(qn("p:cSld"))
    assert cSld is not None
    cSld.set("name", name)


def _build_payments_template_bytes() -> bytes:
    """Build a minimal PPTX with a payments_analysis_stp slide.

    The slide's named placeholders mirror the assembler's expectations
    for ``payments_stp_v1``.
    """
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(blank_layout)
    _set_slide_name(slide, "payments_analysis_stp")

    placeholders = [
        ("title", Inches(0.5), Inches(0.3)),
        ("subtitle", Inches(0.5), Inches(0.9)),
        ("commentary", Inches(0.5), Inches(1.4)),
        ("period_from", Inches(0.5), Inches(2.4)),
        ("period_to", Inches(2.0), Inches(2.4)),
        ("stp_count", Inches(0.5), Inches(2.9)),
        ("non_stp_count", Inches(2.0), Inches(2.9)),
        ("total_count", Inches(3.5), Inches(2.9)),
        ("stp_rate_pct", Inches(0.5), Inches(3.4)),
        ("non_stp_rate_pct", Inches(2.0), Inches(3.4)),
    ]
    width = Inches(1.4)
    height = Inches(0.4)
    for name, left, top in placeholders:
        textbox = slide.shapes.add_textbox(left, top, width, height)
        _name_shape(textbox, name)

    # Reasons table: header + one body row.
    table_shape = slide.shapes.add_table(
        rows=2, cols=3, left=Inches(0.5), top=Inches(4.2),
        width=Inches(8.0), height=Inches(1.0),
    )
    _name_shape(table_shape, "top_reasons_non_stp_table")
    header = table_shape.table.rows[0].cells
    header[0].text = "Reason"
    header[1].text = "Count"
    header[2].text = "%"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _read_named_text(prs: Presentation, slide_index: int, name: str) -> str:
    slide = prs.slides[slide_index]
    for shape in slide.shapes:
        if shape.name == name and shape.has_text_frame:
            return shape.text_frame.text
    raise AssertionError(f"shape {name!r} not found")


def _read_named_table_cell(
    prs: Presentation, slide_index: int, table_name: str, row: int, col: int
) -> str:
    slide = prs.slides[slide_index]
    for shape in slide.shapes:
        if shape.name == table_name and shape.has_table:
            return shape.table.rows[row].cells[col].text
    raise AssertionError(f"table {table_name!r} not found")


# ----------------------------------------------------------------- tests


@pytest.fixture
def template_bytes() -> bytes:
    return _build_payments_template_bytes()


def test_assemble_payments_stp_populates_placeholders(template_bytes: bytes) -> None:
    slide_desc = {
        "slide_id": "payments_stp_v1",
        "title": "Payments STP",
        "subtitle": "Q1 2026",
        "commentary": ["STP rate held above 95% all quarter."],
        "data_gaps": [],
        "model_metadata": {"model": "gpt-4o-mini"},
        "period": {"from": "2026-01-01", "to": "2026-03-31"},
        "volumes": {"stp_count": 9500, "non_stp_count": 500, "total_count": 10000},
        "rates": {"stp_rate_pct": 95.0, "non_stp_rate_pct": 5.0},
        "top_reasons_non_stp": [
            {"reason": "Missing BIC", "count": 200, "pct": 40.0}
        ],
    }

    deck_bytes, failed_slide_ids = DeckAssembler().assemble(
        template_bytes, [slide_desc]
    )

    assert failed_slide_ids == []
    assert deck_bytes
    prs = Presentation(io.BytesIO(deck_bytes))

    assert _read_named_text(prs, 0, "title") == "Payments STP"
    assert _read_named_text(prs, 0, "subtitle") == "Q1 2026"
    assert "STP rate held above 95%" in _read_named_text(prs, 0, "commentary")
    assert _read_named_text(prs, 0, "stp_count") == "9500"
    assert _read_named_text(prs, 0, "non_stp_count") == "500"
    assert _read_named_text(prs, 0, "total_count") == "10000"
    assert _read_named_text(prs, 0, "stp_rate_pct") == "95.0"
    assert _read_named_text(prs, 0, "period_from") == "2026-01-01"

    # Body row populated.
    assert _read_named_table_cell(
        prs, 0, "top_reasons_non_stp_table", row=1, col=0
    ) == "Missing BIC"
    assert _read_named_table_cell(
        prs, 0, "top_reasons_non_stp_table", row=1, col=1
    ) == "200"


def test_assemble_skips_unknown_slide_id_and_records_failure(
    template_bytes: bytes,
) -> None:
    deck_bytes, failed_slide_ids = DeckAssembler().assemble(
        template_bytes,
        [
            {
                "slide_id": "no_such_slide_v1",
                "title": "irrelevant",
                "commentary": [],
                "data_gaps": [],
                "model_metadata": {},
            }
        ],
    )

    assert failed_slide_ids == ["no_such_slide_v1"]
    assert deck_bytes  # template still saved unchanged


def test_assemble_leaves_numeric_fields_blank_when_in_data_gaps(
    template_bytes: bytes,
) -> None:
    slide_desc = {
        "slide_id": "payments_stp_v1",
        "title": "Payments STP",
        "subtitle": None,
        "commentary": [],
        "data_gaps": ["volumes.stp_count", "rates.stp_rate_pct"],
        "model_metadata": {},
        "period": {"from": "2026-01-01", "to": "2026-03-31"},
        "volumes": {"stp_count": None, "non_stp_count": 500, "total_count": 10000},
        "rates": {"stp_rate_pct": None, "non_stp_rate_pct": 5.0},
        "top_reasons_non_stp": [],
    }

    deck_bytes, failed_slide_ids = DeckAssembler().assemble(
        template_bytes, [slide_desc]
    )
    assert failed_slide_ids == []

    prs = Presentation(io.BytesIO(deck_bytes))
    # Numeric fields must stay blank — never carry a marker string.
    assert _read_named_text(prs, 0, "stp_count") == ""
    assert _read_named_text(prs, 0, "stp_rate_pct") == ""
    assert _read_named_text(prs, 0, "non_stp_count") == "500"
