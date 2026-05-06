"""Blob Storage helpers and PPTX deck assembly."""

from __future__ import annotations

import io
from datetime import datetime, timedelta, timezone

from azure.storage.blob import (
    BlobClient,
    BlobSasPermissions,
    BlobServiceClient,
    UserDelegationKey,
    generate_blob_sas,
)
from pptx import Presentation

from shared.auth import get_credential
from shared.config import get_settings
from shared.logging import get_logger
from shared.models.slide import SlideDescription

logger = get_logger(__name__)

_SAS_TTL = timedelta(hours=1)


def _service_client() -> BlobServiceClient:
    settings = get_settings()
    return BlobServiceClient(
        account_url=str(settings.BLOB_ACCOUNT_URL),
        credential=get_credential(),
    )


def _download_template() -> bytes:
    settings = get_settings()
    logger.info(
        "blob.download_template",
        extra={
            "container": settings.BLOB_TEMPLATE_CONTAINER,
            "blob": settings.BLOB_TEMPLATE_NAME,
        },
    )
    with _service_client() as svc:
        blob = svc.get_blob_client(
            container=settings.BLOB_TEMPLATE_CONTAINER,
            blob=settings.BLOB_TEMPLATE_NAME,
        )
        return blob.download_blob().readall()


def _render_pptx(template_bytes: bytes, slides: list[SlideDescription]) -> bytes:
    """Render the deck by injecting each slide's JSON description.

    Phase 1 keeps rendering minimal: the existing template slides are
    preserved and a notes section per slide carries the JSON payload so
    downstream tooling (or a manual reviewer) can verify grounded values.
    Production-quality placeholder mapping is intentionally deferred.
    """
    prs = Presentation(io.BytesIO(template_bytes))

    for idx, slide_desc in enumerate(slides):
        if idx >= len(prs.slides):
            logger.warning(
                "deck.slide_overflow",
                extra={"template_slide_count": len(prs.slides), "incoming_index": idx},
            )
            break
        slide = prs.slides[idx]
        notes_tf = slide.notes_slide.notes_text_frame
        notes_tf.text = slide_desc.model_dump_json()

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _upload_deck(customer_id: str, run_id: str, payload: bytes) -> BlobClient:
    settings = get_settings()
    blob_name = f"{customer_id}/{run_id}.pptx"
    svc = _service_client()
    blob = svc.get_blob_client(
        container=settings.BLOB_DECK_CONTAINER, blob=blob_name
    )
    blob.upload_blob(
        payload,
        overwrite=True,
        content_type=(
            "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        ),
    )
    logger.info(
        "blob.upload_deck",
        extra={
            "container": settings.BLOB_DECK_CONTAINER,
            "blob": blob_name,
            "size_bytes": len(payload),
        },
    )
    return blob


def _sas_url(blob: BlobClient) -> tuple[str, datetime]:
    """Return a user-delegation SAS URL for the deck blob and its expiry."""
    settings = get_settings()
    now = datetime.now(timezone.utc)
    expiry = now + _SAS_TTL

    with _service_client() as svc:
        udk: UserDelegationKey = svc.get_user_delegation_key(
            key_start_time=now - timedelta(minutes=5),
            key_expiry_time=expiry,
        )

    account_name = blob.account_name
    if account_name is None:  # pragma: no cover - defensive
        raise RuntimeError("Blob account_name is not available for SAS generation")

    sas = generate_blob_sas(
        account_name=account_name,
        container_name=blob.container_name,
        blob_name=blob.blob_name,
        user_delegation_key=udk,
        permission=BlobSasPermissions(read=True),
        expiry=expiry,
        start=now - timedelta(minutes=5),
    )
    return f"{blob.url}?{sas}", expiry


def assemble_and_upload(
    *,
    customer_id: str,
    run_id: str,
    slides: list[SlideDescription],
) -> tuple[str, str, datetime | None]:
    """Assemble the deck and persist it to blob storage.

    Returns:
        ``(blob_url, deck_url, expires_at)``. ``deck_url`` is a SAS URL
        when ``DECK_DELIVERY_MODE == 'sas'``; for ``'dataverse'`` mode
        the caller is expected to perform the write-through and resolve
        the record URL — this scaffold returns the bare blob URL.
    """
    settings = get_settings()
    template_bytes = _download_template()
    deck_bytes = _render_pptx(template_bytes, slides)
    blob = _upload_deck(customer_id, run_id, deck_bytes)

    if settings.DECK_DELIVERY_MODE == "sas":
        sas_url, expires_at = _sas_url(blob)
        return blob.url, sas_url, expires_at
    return blob.url, blob.url, None
