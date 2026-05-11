"""Local probe: run _update_chart_xml_cache on slide 21 of latest_deck.pptx."""
from __future__ import annotations
import re, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "asrp_functions"))
from pptx import Presentation
from lxml import etree
from shared.deck_assembler import DeckAssembler

p = Presentation("latest_deck.pptx")
slide = p.slides[20]
for shape in slide.shapes:
    if not shape.has_chart:
        continue
    chart = shape.chart
    DeckAssembler._update_chart_xml_cache(
        chart,
        categories=[f"Cat{i}" for i in range(12)],
        values=[100, 50, 25, 10, 5, 4, 3, 2, 1, 1, 1, 1],
    )
    DeckAssembler._rewrite_data_label_cache(
        chart,
        values=[100, 50, 25, 10, 5, 4, 3, 2, 1, 1, 1, 1],
    )
    xml = etree.tostring(chart._chartSpace, pretty_print=False).decode()
    print("--- a:t after update ---")
    for m in re.findall(r"<a:t>([^<]*)</a:t>", xml):
        print(repr(m))
    print("--- dlblRangeCache snippet ---")
    nc = re.search(r"<c15:dlblRangeCache>.*?</c15:dlblRangeCache>", xml)
    if nc:
        print(nc.group(0)[:600])
    break
